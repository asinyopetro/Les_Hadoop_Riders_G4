# -*- coding: utf-8 -*-
"""PowerPoint G4 — texte enrichi, captures, liens. Conclusion = Wren (dernière)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

TEAL = RGBColor(0x0F, 0x5C, 0x4C)
TEAL_MID = RGBColor(0x14, 0x7A, 0x66)
SLATE = RGBColor(0x1E, 0x29, 0x3B)
GRAY = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF1, 0xF5, 0xF4)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xC2, 0x41, 0x0C)
BLUE = RGBColor(0x1D, 0x4E, 0x89)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
GREEN = RGBColor(0x15, 0x80, 0x3D)

W, H = Inches(13.333), Inches(7.5)
ROOT = Path(__file__).resolve().parents[1]
CAP = ROOT / "captures"

HUB = "https://hub.docker.com/r/leshadoopriders/hadoop-tp-g4"
GITHUB = "https://github.com/asinyopetro/Les_Hadoop_Riders_G4"
UI_HDFS = "http://localhost:9870"
UI_YARN = "http://localhost:8088"
UI_DN = "http://localhost:9870/dfshealth.html#tab-datanode"


def set_run(run, size=20, bold=False, color=SLATE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_hard_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def footer(slide, speaker, page, total=12):
    add_hard_rect(slide, 0, Inches(6.95), W, Inches(0.55), TEAL)
    box = slide.shapes.add_textbox(Inches(0.35), Inches(7.05), Inches(10.2), Inches(0.35))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = f"Présenté par : {speaker}"
    set_run(r, 12, True, WHITE)
    num = slide.shapes.add_textbox(Inches(11.4), Inches(7.05), Inches(1.6), Inches(0.35))
    p2 = num.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = f"{page} / {total}"
    set_run(r2, 11, False, WHITE)


def title_bar(slide, title):
    add_hard_rect(slide, 0, 0, W, Inches(1.05), TEAL)
    add_hard_rect(slide, 0, Inches(1.05), W, Inches(0.07), ACCENT)
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.6))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = title
    set_run(r, 24, True, WHITE)


def textbox(slide, left, top, width, height, lines, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = line.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(line.get("space", 5))
        r = p.add_run()
        r.text = line["text"]
        set_run(r, line.get("size", size), line.get("bold", False), line.get("color", SLATE))
        if line.get("url"):
            r.hyperlink.address = line["url"]
    return box


def link_chip(slide, left, top, width, label, url, fill=TEAL_MID):
    add_rect(slide, Inches(left), Inches(top), Inches(width), Inches(0.4), fill)
    box = slide.shapes.add_textbox(Inches(left), Inches(top + 0.05), Inches(width), Inches(0.32))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    set_run(r, 11, True, WHITE)
    r.hyperlink.address = url


def add_pic(slide, name, left, top, width=None, height=None):
    path = CAP / name
    if not path.exists():
        return None
    kwargs = {}
    if width is not None:
        kwargs["width"] = Inches(width)
    if height is not None:
        kwargs["height"] = Inches(height)
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kwargs)


def caption(slide, left, top, width, text):
    textbox(slide, left, top, width, 0.28, [{"text": text, "size": 11, "color": MUTED}])


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def node_box(slide, left, top, w, h, title, lines, fill=TEAL):
    add_rect(slide, Inches(left), Inches(top), Inches(w), Inches(h), fill)
    textbox(slide, left + 0.1, top + 0.08, w - 0.2, h - 0.15, [
        {"text": title, "size": 13, "bold": True, "color": WHITE, "space": 3},
        *[{"text": x, "size": 11, "color": WHITE, "space": 2} for x in lines],
    ])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # ========== 1 TITRE — Komla ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, TEAL)
    add_hard_rect(s, Inches(8.2), 0, Inches(5.2), H, SLATE)
    add_pic(s, "screenshot-datanodes.png", 8.45, 1.25, width=4.7)
    caption(s, 8.45, 5.35, 4.7, "Preuve : UI NameNode — 5 DataNodes Live")
    textbox(s, 0.55, 1.2, 7.3, 2.2, [
        {"text": "Cluster Hadoop avec Docker", "size": 32, "bold": True, "color": WHITE, "space": 8},
        {"text": "Déploiement et exploitation HDFS + YARN", "size": 18, "color": RGBColor(0xD1, 0xFA, 0xE5), "space": 10},
        {"text": "Cours IFM30522 — UA1 Projet 1", "size": 15, "color": WHITE, "space": 6},
        {"text": "Groupe G4 — Les_Hadoop_Riders", "size": 18, "bold": True, "color": WHITE},
    ])
    textbox(s, 0.55, 3.7, 7.3, 1.2, [
        {"text": "Objectif : montrer qu’on sait déployer un cluster multi-nœuds,", "size": 14, "color": RGBColor(0xE2, 0xE8, 0xF0), "space": 4},
        {"text": "manipuler HDFS, lancer un job YARN et publier l’image Docker.", "size": 14, "color": RGBColor(0xE2, 0xE8, 0xF0)},
    ])
    link_chip(s, 0.55, 5.15, 3.4, "Docker Hub →", HUB)
    link_chip(s, 4.15, 5.15, 3.2, "GitHub →", GITHUB, BLUE)
    footer(s, "Komla Petro Asinyo", 1)
    add_notes(s, "Bonjour, G4 Les_Hadoop_Riders. On présente notre cluster Hadoop Docker. "
              "Objectif en une phrase. Liens Hub et GitHub.")

    # ========== 2 PLAN — Komla ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Plan — qui présente quoi")
    textbox(s, 0.45, 1.2, 12.3, 0.45, [
        {"text": "Environ 8–10 minutes. Chaque slide indique le présentateur en bas.", "size": 14, "color": MUTED},
    ])
    items = [
        ("01", "E-commerce & 3V", "Kassoum"),
        ("02", "HDFS vs FS local", "Joel"),
        ("03", "YARN (RM / NM)", "Forbes"),
        ("04", "Architecture Docker", "Frank"),
        ("05", "Déploiement & Hub", "Frank"),
        ("06", "Manipulations HDFS", "Wren"),
        ("07", "Job π YARN", "Forbes"),
        ("08", "Monitoring UI", "Joel"),
        ("09", "Troubleshooting", "Kassoum"),
        ("10", "Conclusion", "Wren"),
    ]
    # show as two rows of 5 for clarity - wait we have 12 slides mapping differently
    # Keep visual plan matching actual slides 3-12
    plan = [
        ("3", "3V e-commerce", "Kassoum"),
        ("4", "HDFS vs local", "Joel"),
        ("5", "YARN RM / NM", "Forbes"),
        ("6–7", "Archi + Hub", "Frank"),
        ("8", "HDFS pratique", "Wren"),
        ("9", "Job π", "Forbes"),
        ("10", "Monitoring", "Joel"),
        ("11", "Problèmes", "Kassoum"),
        ("12", "Conclusion", "Wren"),
    ]
    for i, (num, label, who) in enumerate(plan):
        col = i % 3
        row = i // 3
        left = 0.45 + col * 4.2
        top = 1.75 + row * 1.55
        add_rect(s, Inches(left), Inches(top), Inches(4.0), Inches(1.4), CARD, TEAL_MID)
        textbox(s, left + 0.2, top + 0.2, 3.6, 1.1, [
            {"text": f"Slide {num}", "size": 12, "bold": True, "color": TEAL, "space": 4},
            {"text": label, "size": 16, "bold": True, "color": SLATE, "space": 4},
            {"text": who, "size": 13, "color": MUTED},
        ])
    footer(s, "Komla Petro Asinyo", 2)
    add_notes(s, "Annoncer l’ordre. Préciser que Wren (dernière) fait la conclusion.")

    # ========== 3 3V — Kassoum ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Cas d’usage : e-commerce et les 3V du Big Data")
    textbox(s, 0.45, 1.2, 12.3, 0.7, [
        {"text": "Nous avons choisi une chaîne de magasins e-commerce : tickets de caisse, logs du site web,", "size": 14, "color": SLATE, "space": 3},
        {"text": "mouvements de stock et images produits. Les données grossissent vite et arrivent en continu.", "size": 14, "color": SLATE},
    ])
    cards = [
        (0.4, "Volume", "Beaucoup d’événements\nsur plusieurs années.\nUne seule machine\nne suffit plus.", BLUE),
        (4.55, "Vélocité", "Commandes et paiements\nen quasi temps réel.\nIl faut pouvoir réagir\n(fraude, rupture).", ORANGE),
        (8.7, "Variété", "Formats mixtes :\nCSV, JSON, logs texte,\nimages produits.", GREEN),
    ]
    for left, title, body, color in cards:
        add_rect(s, Inches(left), Inches(2.05), Inches(3.85), Inches(3.0), CARD, color)
        add_hard_rect(s, Inches(left), Inches(2.05), Inches(3.85), Inches(0.55), color)
        textbox(s, left + 0.2, 2.12, 3.4, 0.45, [{"text": title, "size": 18, "bold": True, "color": WHITE}])
        textbox(s, left + 0.25, 2.8, 3.4, 2.0, [
            {"text": line, "size": 14, "color": SLATE, "space": 5}
            for line in body.split("\n")
        ])
    textbox(s, 0.45, 5.25, 12.3, 1.3, [
        {"text": "Lien avec Hadoop : HDFS stocke ces données de façon distribuée ; YARN permet de lancer", "size": 14, "color": SLATE, "space": 3},
        {"text": "des traitements (ex. MapReduce) sur le cluster sans tout centraliser sur un seul serveur.", "size": 14, "color": SLATE},
    ])
    footer(s, "Kassoum Dene", 3)
    add_notes(s, "Lire le contexte, puis chaque V avec un exemple. Finir sur HDFS + YARN.")

    # ========== 4 HDFS — Joel ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "HDFS vs système de fichiers classique")
    textbox(s, 0.4, 1.18, 12.4, 0.4, [
        {"text": "Différence clé : stockage centralisé sur une machine vs stockage distribué et répliqué.", "size": 14, "color": MUTED},
    ])
    add_rect(s, Inches(0.35), Inches(1.65), Inches(6.15), Inches(4.9), CARD, GRAY)
    textbox(s, 0.55, 1.8, 5.8, 0.4, [{"text": "FS classique (NTFS, ext4…)", "size": 17, "bold": True, "color": GRAY}])
    node_box(s, 1.6, 2.4, 3.5, 1.35, "Une seule machine", ["fichier stocké en entier", "métadonnées gérées par l’OS"], SLATE)
    textbox(s, 0.55, 4.0, 5.8, 2.2, [
        {"text": "• Si le disque tombe, les données sont perdues", "size": 14, "space": 6},
        {"text": "• Pas de réplication automatique entre machines", "size": 14, "space": 6},
        {"text": "• Capacité limitée au matériel local", "size": 14, "space": 6},
        {"text": "• Adapté aux petits volumes du quotidien", "size": 14},
    ])
    add_rect(s, Inches(6.8), Inches(1.65), Inches(6.15), Inches(4.9), CARD, TEAL)
    textbox(s, 7.0, 1.8, 5.8, 0.4, [{"text": "HDFS (Hadoop)", "size": 17, "bold": True, "color": TEAL}])
    node_box(s, 8.2, 2.35, 3.2, 1.05, "NameNode", ["noms, dossiers, emplacement des blocs"], TEAL)
    for i, x in enumerate([7.05, 9.0, 10.95]):
        node_box(s, x, 3.65, 1.8, 1.05, f"DataNode {i+1}", ["blocs répliqués"], TEAL_MID)
    textbox(s, 7.0, 5.0, 5.8, 1.3, [
        {"text": "Fichier découpé en blocs (souvent 128 Mo).", "size": 13, "space": 4},
        {"text": "Plusieurs copies sur le cluster → tolérance aux pannes.", "size": 13, "space": 4},
        {"text": "Le NameNode ne stocke pas le contenu des fichiers.", "size": 13},
    ])
    footer(s, "Joel Kazoni Tugirimana", 4)
    add_notes(s, "Comparer les 2 colonnes. Insister blocs + réplication + rôle NameNode.")

    # ========== 5 YARN — Forbes ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "YARN : qui orchestre, qui exécute ?")
    textbox(s, 0.4, 1.18, 12.4, 0.45, [
        {"text": "YARN (Yet Another Resource Negotiator) gère les ressources CPU/mémoire du cluster pour les jobs.", "size": 14, "color": MUTED},
    ])
    node_box(s, 0.4, 1.75, 3.9, 2.35, "1. Client", ["Soumet une application", "ex. yarn jar … pi 4 1000", "demande des ressources"], BLUE)
    node_box(s, 4.6, 1.75, 4.1, 2.35, "2. ResourceManager", ["Sur le master", "Connaît le cluster", "Alloue des conteneurs", "Démarre l’ApplicationMaster"], TEAL)
    node_box(s, 9.0, 1.75, 3.9, 2.35, "3. NodeManager", ["Sur chaque worker", "Lance les conteneurs", "Surveille l’exécution", "Remonte l’état au RM"], TEAL_MID)
    add_pic(s, "screenshot-yarn.png", 0.4, 4.35, width=7.5)
    caption(s, 0.4, 6.5, 7.5, "Capture UI YARN — applications du cluster")
    textbox(s, 8.2, 4.35, 4.7, 2.2, [
        {"text": "À retenir", "size": 16, "bold": True, "color": TEAL, "space": 8},
        {"text": "ResourceManager = décide et coordonne", "size": 13, "space": 6},
        {"text": "NodeManager = exécute sur sa machine", "size": 13, "space": 10},
        {"text": "Ouvrir UI YARN →", "size": 13, "bold": True, "color": BLUE, "url": UI_YARN},
    ])
    footer(s, "Forbes Magène", 5)
    add_notes(s, "Suivre 1→2→3. Si question ApplicationMaster : chef du job lancé via le RM.")

    # ========== 6 ARCHITECTURE — Frank ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Architecture : 1 master + 5 workers (Docker)")
    textbox(s, 0.4, 1.18, 12.4, 0.4, [
        {"text": "Six conteneurs sur le réseau Docker hadoop-net, construits depuis la même image du groupe.", "size": 14, "color": MUTED},
    ])
    node_box(s, 2.8, 1.7, 7.7, 1.55, "hadoop-master", [
        "NameNode (HDFS)  ·  ResourceManager (YARN)  ·  SecondaryNameNode",
        "Ports exposés : 9870 (UI HDFS) · 8088 (UI YARN) · 9000 (RPC)",
    ], TEAL)
    for i in range(5):
        x = 0.4 + i * 2.55
        node_box(s, x, 3.55, 2.4, 1.35, f"hadoop-worker{i+1}", ["DataNode (stockage)", "NodeManager (jobs)"], TEAL_MID)
    textbox(s, 0.4, 5.15, 12.4, 1.4, [
        {"text": "Image Docker : leshadoopriders/hadoop-tp-g4:1.0   (Hadoop 3.3.6)", "size": 14, "bold": True, "color": SLATE, "space": 6},
        {"text": "UI HDFS → http://localhost:9870", "size": 13, "color": BLUE, "url": UI_HDFS, "space": 3},
        {"text": "UI YARN → http://localhost:8088", "size": 13, "color": BLUE, "url": UI_YARN},
    ])
    footer(s, "Frank A Simo Ngounou", 6)
    add_notes(s, "Schéma master puis 5 workers. Citer ports et image.")

    # ========== 7 DEPLOIEMENT — Frank ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Déploiement du cluster et publication Hub")
    textbox(s, 0.4, 1.2, 6.0, 5.3, [
        {"text": "Mise en route", "size": 16, "bold": True, "color": TEAL, "space": 8},
        {"text": "1. docker compose up --build -d", "size": 14, "space": 5},
        {"text": "   → build local + démarrage des 6 services", "size": 12, "color": MUTED, "space": 8},
        {"text": "2. docker ps", "size": 14, "space": 5},
        {"text": "   → vérifier master + worker1…5", "size": 12, "color": MUTED, "space": 8},
        {"text": "3. hdfs dfsadmin -report", "size": 14, "space": 5},
        {"text": "   → Live datanodes (5) = cluster OK", "size": 12, "color": GREEN, "space": 10},
        {"text": "Publication", "size": 16, "bold": True, "color": TEAL, "space": 8},
        {"text": "docker push leshadoopriders/hadoop-tp-g4:1.0", "size": 13, "space": 6},
        {"text": "Docker Hub impose les minuscules (g4, pas G4).", "size": 12, "color": MUTED, "space": 10},
        {"text": "Page Docker Hub →", "size": 13, "bold": True, "color": BLUE, "url": HUB, "space": 4},
        {"text": "Dépôt GitHub →", "size": 13, "bold": True, "color": BLUE, "url": GITHUB},
    ])
    add_pic(s, "screenshot-terminal-01.png", 6.6, 1.25, width=6.3)
    caption(s, 6.6, 4.85, 6.3, "Terminal : docker ps / dfsadmin -report")
    add_pic(s, "screenshot-namenode.png", 6.6, 5.15, height=1.55)
    footer(s, "Frank A Simo Ngounou", 7)
    add_notes(s, "Dire les 3 commandes. Montrer capture. Rappeler Hub.")

    # ========== 8 HDFS PRATIQUE — Wren ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Manipulations HDFS — dossier /data/ventes/2026")
    textbox(s, 0.35, 1.18, 5.9, 5.4, [
        {"text": "Scénario : données de ventes e-commerce (CSV).", "size": 13, "color": MUTED, "space": 8},
        {"text": "1. Création du dossier + droits", "size": 14, "bold": True, "color": TEAL, "space": 3},
        {"text": "mkdir -p /data/ventes/2026  puis chmod 755", "size": 12, "space": 7},
        {"text": "2. Upload du fichier", "size": 14, "bold": True, "color": TEAL, "space": 3},
        {"text": "put transactions.csv  + chmod 644", "size": 12, "space": 7},
        {"text": "3. Vérification des blocs", "size": 14, "bold": True, "color": TEAL, "space": 3},
        {"text": "fsck + stat → HEALTHY, 1 bloc, réplication 3", "size": 12, "space": 7},
        {"text": "4. Lecture et fusion", "size": 14, "bold": True, "color": TEAL, "space": 3},
        {"text": "cat | head  puis  getmerge de 2 CSV", "size": 12, "space": 7},
        {"text": "5. Corbeille + setrep", "size": 14, "bold": True, "color": TEAL, "space": 3},
        {"text": "rm vers .Trash  ;  setrep -w 2 → réplication 2", "size": 12, "color": GREEN},
    ])
    add_pic(s, "screenshot-terminal-02.png", 6.4, 1.25, width=6.5)
    caption(s, 6.4, 5.35, 6.5, "Capture terminal — commandes HDFS réellement exécutées")
    textbox(s, 6.4, 5.65, 6.5, 0.9, [
        {"text": "Note : fichier petit (377 o) → 1 seul bloc. Normal avec une taille de bloc de 128 Mo.", "size": 12, "color": MUTED},
    ])
    footer(s, "Wren Surprenant-Nicolson", 8)
    add_notes(s, "Suivre les 5 étapes. Expliquer Trash et setrep. Reviendra pour la conclusion.")

    # ========== 9 JOB PI — Forbes ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Job YARN : calcul de π (exemple MapReduce)")
    textbox(s, 0.35, 1.18, 5.8, 5.4, [
        {"text": "Nous avons lancé l’exemple officiel Hadoop (Monte Carlo)", "size": 13, "color": MUTED, "space": 8},
        {"text": "Commande", "size": 14, "bold": True, "color": TEAL, "space": 4},
        {"text": "yarn jar …/hadoop-mapreduce-examples-3.3.6.jar pi 4 1000", "size": 12, "space": 4},
        {"text": "→ 4 maps, 1000 samples par map", "size": 12, "color": MUTED, "space": 10},
        {"text": "Résultat obtenu", "size": 14, "bold": True, "color": TEAL, "space": 6},
        {"text": "Application : application_1791039798429_0001", "size": 13, "space": 5},
        {"text": "Nom : QuasiMonteCarlo", "size": 13, "space": 5},
        {"text": "État final : SUCCEEDED", "size": 16, "bold": True, "color": GREEN, "space": 5},
        {"text": "Estimation : π ≈ 3.14", "size": 15, "bold": True, "color": SLATE, "space": 10},
        {"text": "Cela prouve que YARN alloue bien des ressources", "size": 12, "space": 3},
        {"text": "et que les NodeManagers exécutent le job.", "size": 12, "space": 8},
        {"text": "Ouvrir UI YARN →", "size": 13, "bold": True, "color": BLUE, "url": UI_YARN},
    ])
    add_pic(s, "screenshot-yarn-app.png", 6.3, 1.2, width=6.6)
    caption(s, 6.3, 4.7, 6.6, "UI — détail application FINISHED / SUCCEEDED")
    add_pic(s, "screenshot-terminal-03.png", 6.3, 5.0, height=1.7)
    footer(s, "Forbes Magène", 9)
    add_notes(s, "Expliquer vite Monte Carlo. Insister SUCCEEDED + id application.")

    # ========== 10 MONITORING — Joel ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Monitoring : interfaces web NameNode et YARN")
    textbox(s, 0.35, 1.15, 12.5, 0.4, [
        {"text": "Les UI permettent de vérifier l’état du cluster sans rester uniquement en ligne de commande.", "size": 13, "color": MUTED},
    ])
    add_pic(s, "screenshot-namenode.png", 0.35, 1.6, width=6.15)
    caption(s, 0.35, 4.55, 6.15, "NameNode — vue d’ensemble du DFS")
    add_pic(s, "screenshot-datanodes.png", 6.8, 1.6, width=6.1)
    caption(s, 6.8, 4.55, 6.1, "Onglet DataNodes — 5 nœuds actifs")
    textbox(s, 0.35, 4.85, 12.5, 1.7, [
        {"text": "NameNode (9870) : 5 DataNodes Live, espace DFS faible (labo), aucun missing block pendant nos tests.", "size": 13, "space": 4},
        {"text": "YARN (8088) : job π en FINISHED/SUCCEEDED ; métriques mémoire et vcores visibles.", "size": 13, "space": 4},
        {"text": "Après setrep -w 2, le NameNode met à jour la cible de réplication du fichier.", "size": 13, "space": 6},
        {"text": "NameNode →", "size": 12, "bold": True, "color": BLUE, "url": UI_HDFS, "space": 2},
        {"text": "DataNodes →", "size": 12, "bold": True, "color": BLUE, "url": UI_DN, "space": 2},
        {"text": "YARN →", "size": 12, "bold": True, "color": BLUE, "url": UI_YARN},
    ])
    footer(s, "Joel Kazoni Tugirimana", 10)
    add_notes(s, "Montrer les 2 captures. Citer missing block / setrep.")

    # ========== 11 TROUBLESHOOTING — Kassoum ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Problèmes rencontrés et solutions")
    textbox(s, 0.4, 1.15, 12.4, 0.4, [
        {"text": "Trois blocages réels pendant le projet — et comment on les a corrigés.", "size": 14, "color": MUTED},
    ])
    problems = [
        ("1. pull access denied",
         "Compose essayait de télécharger l’image du groupe avant qu’elle existe sur Hub.",
         "pull_policy: never + build local de l’image, puis push une fois prête.", ACCENT),
        ("2. Tag avec majuscules refusé",
         "Docker Hub n’accepte pas les majuscules dans le nom de dépôt (ex. G4).",
         "Tag final : leshadoopriders/hadoop-tp-g4:1.0 (tout en minuscules).", ORANGE),
        ("3. Téléchargement Hadoop trop lent",
         "Le build restait bloqué longtemps sur archive.apache.org.",
         "Image de base apache/hadoop:3.3.6 + notre config XML et entrypoint.sh.", BLUE),
    ]
    for i, (title, cause, sol, color) in enumerate(problems):
        top = 1.6 + i * 1.65
        add_rect(s, Inches(0.4), Inches(top), Inches(12.5), Inches(1.5), CARD, color)
        add_hard_rect(s, Inches(0.4), Inches(top), Inches(0.16), Inches(1.5), color)
        textbox(s, 0.75, top + 0.12, 11.9, 1.3, [
            {"text": title, "size": 15, "bold": True, "color": SLATE, "space": 3},
            {"text": f"Cause : {cause}", "size": 12, "color": MUTED, "space": 3},
            {"text": f"Solution : {sol}", "size": 13, "bold": True, "color": TEAL},
        ])
    footer(s, "Kassoum Dene", 11)
    add_notes(s, "Un problème = une cause = une solution. Puis passer à Wren pour conclure.")

    # ========== 12 CONCLUSION — Wren (dernière) ==========
    s = prs.slides.add_slide(blank)
    add_hard_rect(s, 0, 0, W, H, CREAM)
    title_bar(s, "Conclusion — ce que nous avons livré")
    textbox(s, 0.4, 1.15, 7.9, 0.45, [
        {"text": "En résumé, le groupe a déployé, testé et documenté un cluster Hadoop complet.", "size": 14, "color": MUTED},
    ])
    points = [
        "Cluster opérationnel : 1 master + 5 DataNodes / NodeManagers",
        "HDFS maîtrisé : droits, fsck, lecture, getmerge, Trash, setrep",
        "Job MapReduce π terminé avec succès (SUCCEEDED, π ≈ 3.14)",
        "Image publiée : leshadoopriders/hadoop-tp-g4:1.0",
        "Rapport + code source + captures pour le rendu",
    ]
    for i, b in enumerate(points):
        y = 1.65 + i * 0.72
        add_rect(s, Inches(0.4), Inches(y), Inches(7.9), Inches(0.62), CARD, TEAL_MID)
        textbox(s, 0.6, y + 0.14, 7.5, 0.4, [{"text": f"✓  {b}", "size": 13, "bold": True, "color": SLATE}])
    add_pic(s, "screenshot-yarn-app.png", 8.55, 1.55, width=4.4)
    caption(s, 8.55, 4.35, 4.4, "Preuve job SUCCEEDED")
    link_chip(s, 8.55, 4.7, 4.4, "Docker Hub →", HUB)
    link_chip(s, 8.55, 5.25, 4.4, "GitHub →", GITHUB, BLUE)
    textbox(s, 0.4, 5.4, 7.9, 1.2, [
        {"text": "Merci — des questions ?", "size": 26, "bold": True, "color": TEAL, "space": 6},
        {"text": "Les_Hadoop_Riders  ·  Komla · Kassoum · Joel · Forbes · Frank · Wren", "size": 12, "color": MUTED},
    ])
    footer(s, "Wren Surprenant-Nicolson", 12)
    add_notes(s, "Wren conclut : 5 points rapidement, liens Hub/GitHub, puis « Merci, questions ? ».")

    out = ROOT / "presentation" / "PRESENTATION-G4-Les_Hadoop_Riders.pptx"
    prs.save(str(out))
    print("OK", out)
    return out


if __name__ == "__main__":
    build()
